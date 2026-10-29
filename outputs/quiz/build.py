from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

root=Path(__file__).parent
src=Document(r'C:/Users/황석준/Desktop/과제 폴더/2학년-2/언어속 문화읽기/퀴즈 중간고사.docx')
t=[p.text.strip() for p in src.paragraphs]
d=Document()
s=d.sections[0]
s.page_width=Cm(21); s.page_height=Cm(29.7)
s.top_margin=Cm(1.8); s.bottom_margin=Cm(1.8)
s.left_margin=Cm(2); s.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name]; st.font.name='맑은 고딕'; st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
 st.font.size=Pt(10.5)
 st.paragraph_format.line_spacing=1.18
 st.paragraph_format.space_after=Pt(4)
d.styles['Title'].font.size=Pt(20)
d.styles['Heading 1'].font.size=Pt(13)
d.styles['Heading 1'].font.bold=True
d.styles['Heading 1'].paragraph_format.space_before=Pt(13)
d.styles['Heading 1'].paragraph_format.space_after=Pt(7)
d.add_paragraph('언어 속 문화읽기 중간고사 퀴즈', 'Title')
d.add_paragraph('9개 팀 · 총 13문제', 'Normal')
num=0
def team(n,title,new=False):
 p=d.add_paragraph(f'{n}팀  {title}','Heading 1')
 if new:p.paragraph_format.page_break_before=True
def para(text,bold=False,keep=True):
 p=d.add_paragraph(); p.paragraph_format.keep_with_next=keep
 r=p.add_run(text); r.bold=bold
 return p
def question(text,options,answer,explain=None):
 global num
 num+=1
 para(f'문제 {num:02d}  {text}',True)
 for i,option in enumerate(options):
  p=para(f'{i+1}. {option}')
  p.paragraph_format.left_indent=Cm(.35)
  p.paragraph_format.space_after=Pt(2)
 p=para('정답  '+answer,True,explain is not None)
 p.runs[0].font.color.rgb=RGBColor.from_string('17568A')
 p.paragraph_format.space_before=Pt(4)
 p.paragraph_format.space_after=Pt(9 if not explain else 3)
 if explain:
  p=para('해설  '+explain,False,False)
  p.paragraph_format.space_after=Pt(9)
team(1,'자연환경과 음식')
question('각 나라의 음식들 중 자연환경의 영향을 받아 발달한 음식을 나라별로 하나씩 쓰면?',[], '아이슬란드 — 루그브라우드 / 스위스 — 치즈퐁뒤 / 일본 — 스시')
team(2,'게임과 인터넷카페 문화')
question('일본 게임 속 몬스터 디자인이 빙산의 표면적 문화라면, 그 밑에 깔려 있는 심층적인 문화는 무엇일까요?', [x[1:] for x in t[6:10]],'② 만물에 영혼이 깃들어 있다고 믿는 애니미즘')
question('일본 인터넷카페 문화에서 나타난 대표적인 사회현상 중 하나로, 일정한 주거지를 확보하지 못한 일부 사람들이 인터넷 카페 등을 숙박 장소로 이용하는 현상은?',[],'넷카페 난민')
question('한국인이 게임을 잘하는 이유와 배경으로 알맞지 않은 것은 무엇일까요?',t[12:16],'③ 자유롭고 여유로운 낭만을 추구하는 플레이 성향')
team(3,'일본의 신앙과 불교',True)
question('일본의 토착적인 가미 신앙이 불교가 전래된 이후 사라지지 않고 불교와 결합한 현상을 신불습합이라 한다. 이 현상 이후 부처와 가미의 관계를 설명하기 위해 등장한 사상을 무엇이라고 하는가?',[],'본지수적설')
team(4,'일본의 식사 예절')
question('일본에서 젓가락끼리 음식을 건네지 않는 이유는?', [x[2:] for x in t[18].splitlines()],'② 장례 의례를 연상시키기 때문에',t[19]+'.')
team(5,'종교와 신앙')
question('다음 두 질문에 답하시오.',[], '(1) 카르마(karma)  /  (2) 마조')
# Put the two subquestions before this question’s answer.
answer=d.paragraphs[-1]
for line in ['(1) 현생에서의 행위가 다음 생의 계급, 신분 등을 결정하는 용어는?', '(2) 대만의 도교에서 주신으로 모시는 신은?']:
 p=answer.insert_paragraph_before(line); p.paragraph_format.keep_with_next=True
team(6,'세계의 팁과 인사 매너',True)
question('미국의 팁 문화가 오늘날과 같이 자리 잡게 된 배경이 아닌 것은?',t[24:28],'④ 미국 정부가 관광산업을 위해 국가적 제도 측면에서 통제한 것')
question('프랑스에 여행을 간 학생이 처음 방문한 작은 가게에 들어갔습니다. 이때 가장 적절한 행동은 무엇일까요?', [x[3:] for x in t[30].splitlines()], '④ 직원에게 먼저 ‘Bonjour’라고 인사한다.')
team(7,'미국 문학과 아메리칸 드림')
question('다음 설명에 해당하는 작품은 무엇일까요?\n'+t[32].removesuffix(' 2번'),[x[2:] for x in t[33].splitlines()],'② 《위대한 개츠비》 — F. 스콧 피츠제럴드','《위대한 개츠비》는 1920년대 미국의 경제 호황기(재즈 시대)를 배경으로 부와 성공을 꿈꾸는 개츠비의 비극을 그린 작품이다. 미국 문학의 핵심 주제인 ‘아메리칸 드림’의 빛과 그림자를 동시에 보여주며, 자유롭게 꿈을 이루고자 하는 열망과 그 허상을 비판적으로 다룬다.')
team(8,'아일랜드의 전통 스포츠',True)
question('이번 발표에서 살펴본 아일랜드의 전통 스포츠는 사람들에게 어떤 의미를 갖는가?',t[36:40], '② 지역의 소속감을 나누고 고유한 문화와 정체성을 이어가는 활동이다.')
team(9,'J-POP과 K-POP')
question('다음 중 J-POP의 특징과 그에 연결되는 일본 문화에 대한 설명으로 옳은 것은?',t[42:46], '② J-POP에서는 음악과 애니메이션·영화 등의 이야기를 함께 경험하는 문화가 나타난다.')
question('K-POP의 특징으로 옳지 않은 것은?',t[47:51], '① K-POP이라는 용어는 1990년대 한국에서 처음 만들어져 사용되기 시작했다.')
assert num==13
for p in d.paragraphs:
 p.paragraph_format.widow_control=True
 p.paragraph_format.keep_together=True
footer=s.footer.paragraphs[0]; footer.alignment=2
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
d.save(root/'퀴즈 중간고사 정리본.docx')
print('Created 9 team sections and',num,'questions')
